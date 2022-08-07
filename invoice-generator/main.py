import os

from docx import Document
from openpyxl import load_workbook


def generate_invoice(invoice_data, path, file_number):
    document = Document(template)
    table = document.tables[0]

    # fill up invoice
    table.rows[3].cells[2].text = str(invoice_data[7]) + "(" + str(invoice_data[0]) + ")"  # phone + customer
    table.rows[5].cells[2].text = str(invoice_data[2]) + str(invoice_data[6])   # street adr + postcode
    table.rows[6].cells[2].text = str(invoice_data[8])  # email

    # invoice name
    document.save(path + "Invoice_" + str(file_number) + ".docx")


if __name__ == "__main__":
    # paths
    cwd = os.getcwd()
    static_path = cwd + "\\static\\data\\"
    data = static_path + "details.xlsx"
    template = static_path + "template.docx"
    generate = static_path + "\\generated_files\\"

    # init workbook
    wb = load_workbook(data)
    ws = wb.active

    # loop thru excel data
    for count, row in enumerate(ws.values):
        row_data = []
        print(row)

        for cell in row:
            # skip header
            if count == 0:
                continue

            else:
                if cell is None:
                    row_data.append("")
                else:
                    row_data.append(cell)

        # if not header, generate invoice
        if count != 0:
            generate_invoice(row_data, generate, count)
